import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl


ROOT = Path.cwd()
ATLAS_PATH = ROOT / "public" / "data" / "atlas.json"

SHEET_NAME = "НПЗ свежие реестры"


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


def number(value):
    """Преобразует Excel-значение в число, если возможно."""
    value = clean(value)

    if value is None:
        return None

    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip().replace(" ", "").replace(",", ".")

    try:
        return float(text)
    except ValueError:
        return None


def json_safe(value):
    if isinstance(value, (datetime,)):
        return value.isoformat()

    if isinstance(value, dict):
        return {str(k): json_safe(v) for k, v in value.items()}

    if isinstance(value, (list, tuple)):
        return [json_safe(v) for v in value]

    return value


def normalize_country(value):
    if value is None:
        return ""

    text = str(value).strip().casefold()

    aliases = {
        "united states of america": "united states",
        "usa": "united states",
        "u.s.a.": "united states",
        "uk": "united kingdom",
        "czechia": "czech republic",
    }

    return aliases.get(text, text)


def map_status(source_status):
    """
    Приводим статусы свежего реестра к формату приложения.
    Исходный статус при этом полностью сохраняется в raw.
    """
    s = str(source_status or "").strip().casefold()

    if s in {
        "open",
        "operable",
        "operating",
        "installed",
        "capacity augmentation",
    }:
        return "Active"

    if s == "closed":
        return "Closed"

    return "Unknown"


if len(sys.argv) < 2:
    print("Использование:")
    print('py scripts\\update_refineries.py "Страны_и_НПЗ_обновлено_2026-09-16.xlsx"')
    sys.exit(1)


XLSX_PATH = Path(sys.argv[1])

if not XLSX_PATH.is_absolute():
    XLSX_PATH = ROOT / XLSX_PATH


if not XLSX_PATH.exists():
    print(f"ОШИБКА: Excel не найден:")
    print(XLSX_PATH)
    sys.exit(1)


if not ATLAS_PATH.exists():
    print(f"ОШИБКА: atlas.json не найден:")
    print(ATLAS_PATH)
    sys.exit(1)


print("Excel:", XLSX_PATH)
print("Atlas:", ATLAS_PATH)
print()


# ---------------------------------------------------------
# Читаем atlas.json
# ---------------------------------------------------------

with ATLAS_PATH.open("r", encoding="utf-8") as f:
    atlas = json.load(f)


countries = atlas.get("countries", [])
old_refineries = atlas.get("refineries", [])


if not isinstance(countries, list):
    raise RuntimeError("atlas['countries'] должен быть списком")

if not isinstance(old_refineries, list):
    raise RuntimeError("atlas['refineries'] должен быть списком")


print("Старых НПЗ в atlas:", len(old_refineries))


# ---------------------------------------------------------
# Строим справочник стран
# ---------------------------------------------------------

country_lookup = {}

for country in countries:
    cid = str(country.get("id"))

    candidates = [
        country.get("name"),
        country.get("nameRu"),
        (country.get("raw") or {}).get("NAME_EN"),
        (country.get("raw") or {}).get("NAME_RU"),
    ]

    for candidate in candidates:
        if candidate:
            country_lookup[normalize_country(candidate)] = cid


# Дополнительные соответствия
country_aliases = {
    "united states of america": "united states",
    "usa": "united states",
    "u.s.a.": "united states",
    "czechia": "czech republic",
    "uk": "united kingdom",
}

for alias, canonical in country_aliases.items():
    canonical_id = country_lookup.get(normalize_country(canonical))
    if canonical_id:
        country_lookup[normalize_country(alias)] = canonical_id


# ---------------------------------------------------------
# Читаем Excel
# ---------------------------------------------------------

wb = openpyxl.load_workbook(
    XLSX_PATH,
    read_only=True,
    data_only=True
)

if SHEET_NAME not in wb.sheetnames:
    print("ОШИБКА: лист не найден:", SHEET_NAME)
    print("Доступные листы:")
    for s in wb.sheetnames:
        print(" -", s)
    sys.exit(1)


ws = wb[SHEET_NAME]

iterator = ws.iter_rows(values_only=True)

try:
    headers = [clean(v) for v in next(iterator)]
except StopIteration:
    raise RuntimeError("Лист НПЗ пуст")


required = {
    "Страна",
    "НПЗ / площадка",
    "Компания",
    "Статус источника",
    "Мощность, барр./сут.",
    "Мощность, млн т/год",
    "x",
    "y",
}

missing_columns = required - set(headers)

if missing_columns:
    print("ОШИБКА: отсутствуют колонки:")
    for c in sorted(missing_columns):
        print(" -", c)
    sys.exit(1)


excel_rows = []

for values in iterator:
    row = dict(zip(headers, values))

    if not any(v is not None for v in values):
        continue

    if clean(row.get("Страна")) is None:
        continue

    if clean(row.get("НПЗ / площадка")) is None:
        continue

    excel_rows.append(row)


print("Строк НПЗ в свежем реестре:", len(excel_rows))


# ---------------------------------------------------------
# Определяем страны, которые реально обновляем
# ---------------------------------------------------------

fresh_country_ids = set()
unmatched_countries = set()

for row in excel_rows:
    country_name = clean(row.get("Страна"))
    key = normalize_country(country_name)

    country_id = country_lookup.get(key)

    if country_id is None:
        unmatched_countries.add(str(country_name))
    else:
        fresh_country_ids.add(str(country_id))


print("Стран со свежим реестром:", len(fresh_country_ids))

if unmatched_countries:
    print()
    print("ВНИМАНИЕ: не сопоставлены страны:")
    for name in sorted(unmatched_countries):
        print(" -", name)


# ---------------------------------------------------------
# Сохраняем старые НПЗ стран, которых нет в свежем реестре
# ---------------------------------------------------------

preserved_refineries = []

removed_old = 0

for refinery in old_refineries:
    country_id = str(refinery.get("country"))

    if country_id in fresh_country_ids:
        removed_old += 1
    else:
        preserved_refineries.append(refinery)


print()
print("Старых НПЗ заменяется:", removed_old)
print("Старых НПЗ сохраняется:", len(preserved_refineries))


# ---------------------------------------------------------
# ID новых НПЗ
# ---------------------------------------------------------

existing_ids = []

for r in preserved_refineries:
    try:
        existing_ids.append(int(r.get("id")))
    except (TypeError, ValueError):
        pass


next_id = max(existing_ids, default=0) + 1


# ---------------------------------------------------------
# Создаём новые записи НПЗ
# ---------------------------------------------------------

new_refineries = []

invalid_coordinates = 0
skipped_country = 0


for row in excel_rows:

    country_name = clean(row.get("Страна"))
    country_id = country_lookup.get(
        normalize_country(country_name)
    )

    if country_id is None:
        skipped_country += 1
        continue


    name = clean(row.get("НПЗ / площадка"))
    company = clean(row.get("Компания"))

    source_status = clean(row.get("Статус источника"))
    status = map_status(source_status)

    x = number(row.get("x"))
    y = number(row.get("y"))

    valid_coordinates = (
        x is not None
        and y is not None
        and -180 <= x <= 180
        and -90 <= y <= 90
    )

    if not valid_coordinates:
        invalid_coordinates += 1


    capacity_bpd = number(
        row.get("Мощность, барр./сут.")
    )

    capacity_million_t = number(
        row.get("Мощность, млн т/год")
    )

    capacity_tonnes = None

    if capacity_million_t is not None:
        capacity_tonnes = capacity_million_t * 1_000_000


    # Сохраняем абсолютно все исходные колонки Excel
    raw = {}

    for key, value in row.items():
        raw[str(key)] = json_safe(clean(value))


    # Добавляем совместимые поля со старой структурой
    raw["Официальное название"] = name
    raw["Владелец завода"] = company

    if capacity_tonnes is not None:
        raw["Мощность переработки, т/год"] = capacity_tonnes

    raw["Статус"] = source_status
    raw["Статус на 2026 г."] = source_status

    if clean(row.get("URL")):
        raw["Источники (URL)"] = clean(row.get("URL"))


    refinery = {
        "id": str(next_id),

        "name": name,

        "originalName": name,

        "coordinates": [
            x,
            y
        ] if valid_coordinates else [None, None],

        "capacity": capacity_tonnes,

        "capacityBpd": capacity_bpd,

        "owner": company,

        "status": status,

        "rawStatus": source_status,

        "age": None,

        "statusConflict": False,

        "raw": raw,

        "validCoordinates": valid_coordinates,

        "country": str(country_id),

        "method": "fresh_registry",

        "distanceKm": 0,
    }


    new_refineries.append(refinery)

    next_id += 1


# ---------------------------------------------------------
# Объединяем старые и новые
# ---------------------------------------------------------

all_refineries = (
    preserved_refineries
    + new_refineries
)


atlas["refineries"] = all_refineries


# ---------------------------------------------------------
# Обновляем показатели НПЗ внутри countries
# ---------------------------------------------------------

by_country = {}

for refinery in all_refineries:

    cid = str(refinery.get("country"))

    if cid not in by_country:
        by_country[cid] = []

    by_country[cid].append(refinery)


for country in countries:

    cid = str(country.get("id"))

    # Меняем статистику только для стран свежего реестра
    if cid not in fresh_country_ids:
        continue

    refs = by_country.get(cid, [])

    active_refs = [
        r for r in refs
        if r.get("status") == "Active"
    ]

    known_capacity = [
        r.get("capacityBpd")
        for r in active_refs
        if isinstance(r.get("capacityBpd"), (int, float))
    ]

    raw = country.setdefault("raw", {})

    raw["НПЗ Oilmap, шт."] = len(active_refs)

    raw["Известная мощность Oilmap, барр./сут."] = sum(
        known_capacity
    )

    raw["НПЗ Oilmap без мощности, шт."] = sum(
        1
        for r in active_refs
        if r.get("capacityBpd") is None
    )

    raw["Дата слоя Oilmap"] = "2026-09-16"

    raw["НПЗ: статус свежей проверки"] = (
        "Обновлено из листа «НПЗ свежие реестры»"
    )


# ---------------------------------------------------------
# Quality
# ---------------------------------------------------------

quality = atlas.setdefault("quality", {})

quality["refineryRecords"] = len(all_refineries)
quality["freshRefineryRecords"] = len(new_refineries)
quality["freshRefineryCountries"] = len(fresh_country_ids)


# ---------------------------------------------------------
# Backup
# ---------------------------------------------------------

timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

backup_path = (
    ATLAS_PATH.parent
    / f"atlas.backup_before_refineries_{timestamp}.json"
)

shutil.copy2(
    ATLAS_PATH,
    backup_path
)


# ---------------------------------------------------------
# Записываем atlas.json
# ---------------------------------------------------------

with ATLAS_PATH.open(
    "w",
    encoding="utf-8"
) as f:

    json.dump(
        atlas,
        f,
        ensure_ascii=False,
        separators=(",", ":"),
        allow_nan=False,
    )


print()
print("=" * 60)
print("ГОТОВО")
print("=" * 60)

print("Backup:")
print(backup_path)

print()
print("Старых НПЗ было:", len(old_refineries))
print("Старых НПЗ заменено:", removed_old)
print("Старых НПЗ сохранено:", len(preserved_refineries))
print("Свежих НПЗ добавлено:", len(new_refineries))
print("Всего НПЗ теперь:", len(all_refineries))
print("Стран обновлено:", len(fresh_country_ids))
print("НПЗ с некорректными координатами:", invalid_coordinates)
print("Строк пропущено из-за страны:", skipped_country)

if unmatched_countries:
    print()
    print("Не сопоставленные страны:")
    for country in sorted(unmatched_countries):
        print(" -", country)

print()
print("Теперь скопируйте atlas.json в dist.")