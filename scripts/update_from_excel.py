import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl


ROOT = Path.cwd()
ATLAS_PATH = ROOT / "public" / "data" / "atlas.json"


def to_number(value):
    if value is None or value == "" or value == "н/д":
        return None

    if isinstance(value, (int, float)):
        return value

    try:
        cleaned = str(value).replace(" ", "").replace(",", ".")
        return float(cleaned)
    except (ValueError, TypeError):
        return None


if len(sys.argv) < 2:
    print("Укажите Excel-файл.")
    print(
        'Пример: py scripts\\update_from_excel.py '
        '"Страны_и_НПЗ_обновлено_2026-09-16.xlsx"'
    )
    sys.exit(1)


xlsx_path = Path(sys.argv[1])

if not xlsx_path.exists():
    print(f"Excel-файл не найден: {xlsx_path}")
    sys.exit(1)

if not ATLAS_PATH.exists():
    print(f"atlas.json не найден: {ATLAS_PATH}")
    sys.exit(1)


print("Читаю Excel...")

wb = openpyxl.load_workbook(
    xlsx_path,
    read_only=True,
    data_only=True,
)

sheet_name = "country_world_globalpetrolprice"

if sheet_name not in wb.sheetnames:
    print(f"Не найден лист: {sheet_name}")
    print("Доступные листы:")
    for name in wb.sheetnames:
        print(f"  - {name}")
    sys.exit(1)

ws = wb[sheet_name]

rows = ws.iter_rows(values_only=True)

try:
    headers = next(rows)
except StopIteration:
    print("Лист Excel пуст.")
    sys.exit(1)

excel_countries = []

for row in rows:
    obj = {}

    for key, value in zip(headers, row):
        if key is not None:
            obj[str(key)] = value

    if obj.get("fid") is not None:
        excel_countries.append(obj)


print(
    f"Excel: найдено стран/территорий: "
    f"{len(excel_countries)}"
)


print("Читаю atlas.json...")

with open(ATLAS_PATH, "r", encoding="utf-8") as f:
    atlas = json.load(f)


countries = atlas.get("countries")

if not isinstance(countries, list):
    print("Ошибка: atlas.json не содержит массив countries.")
    sys.exit(1)


FIELDS = {
    "consumption":
        "Потребление нефти/нефтепродуктов, барр./сут.",

    "production":
        "Добыча нефти и жидкостей, барр./сут.",

    "crude":
        "Добыча сырой нефти, барр./сут.",

    "ngpl":
        "Добыча NGPL, барр./сут.",

    "gdp":
        "ВВП, млрд USD",

    "growth":
        "Рост ВВП, %",

    "inflation":
        "Инфляция CPI, %",

    "diesel":
        "Дизель/дистилляты, % потребления",

    "gasoline":
        "Бензин, % потребления",

    "lpg":
        "LPG, % потребления",

    "jet":
        "Jet/керосин, % потребления",

    "residual":
        "Мазут/Residual Fuel Oil, % потребления",

    "gasolinePrice":
        "Цена бензина, USD",

    "dieselPrice":
        "Цена дизеля, USD",

    "jetPrice":
        "Цена Jet/керосина, USD",

    "lpgPrice":
        "Цена LPG, USD",

    "residualPrice":
        "Цена мазута, USD",
}


YEAR_FIELDS = {
    "consumption": "Год потребления",
    "production": "Год добычи",
    "crude": "Год сырой нефти / конденсата",
    "ngpl": "Год NGPL",

    "gdp": "Год ВВП",
    "growth": "Год роста",
    "inflation": "Год инфляции",

    "diesel": "Год продуктовых долей",
    "gasoline": "Год продуктовых долей",
    "lpg": "Год продуктовых долей",
    "jet": "Год продуктовых долей",
    "residual": "Год продуктовых долей",
}


def normalize_name(value):
    if value is None:
        return ""

    return str(value).strip().casefold()


# Индекс стран по ID
by_id = {}

for country in countries:
    country_id = country.get("id")

    if country_id is not None:
        by_id[str(country_id)] = country


# Индекс по английскому названию
by_name = {}

for country in countries:
    name = normalize_name(country.get("name"))

    if name:
        by_name[name] = country


updated = 0
missing = []


print("Обновляю страны...")


for row in excel_countries:

    fid = str(row.get("fid"))

    country = by_id.get(fid)

    # Если fid не совпал — пробуем название
    if country is None:
        excel_name = normalize_name(
            row.get("NAME_EN")
        )

        country = by_name.get(excel_name)


    if country is None:
        missing.append({
            "fid": fid,
            "name": row.get("NAME_EN"),
        })
        continue


    # Сохраняем исходные поля Excel
    raw = country.get("raw")

    if not isinstance(raw, dict):
        raw = {}

    for key, value in row.items():

        # datetime нельзя напрямую записать в JSON
        if isinstance(value, datetime):
            raw[key] = value.isoformat()

        else:
            raw[key] = value

    country["raw"] = raw


    # Названия
    if row.get("NAME_EN"):
        country["name"] = row["NAME_EN"]

    if row.get("NAME_RU"):
        country["nameRu"] = row["NAME_RU"]


    # Числовые показатели
    values = country.get("values")

    if not isinstance(values, dict):
        values = {}

    for key, excel_field in FIELDS.items():

        value = to_number(
            row.get(excel_field)
        )

        # Пустые значения Excel не уничтожают
        # существующие данные.
        if value is not None:
            values[key] = value

    country["values"] = values


    # Годы
    years = country.get("years")

    if not isinstance(years, dict):
        years = {}

    for key, excel_field in YEAR_FIELDS.items():

        value = to_number(
            row.get(excel_field)
        )

        if value is not None:

            if isinstance(value, float) and value.is_integer():
                value = int(value)

            years[key] = value

    country["years"] = years


    # Информация о ценах
    if row.get("Единица/тип цены"):
        country["priceUnit"] = (
            row["Единица/тип цены"]
        )

    if row.get("Дата цены") is not None:

        price_date = row["Дата цены"]

        if isinstance(price_date, datetime):
            price_date = price_date.isoformat()

        country["priceDate"] = price_date


    updated += 1


# Обновляем quality
quality = atlas.get("quality")

if not isinstance(quality, dict):
    quality = {}

quality["countryRecords"] = len(countries)


price_records = 0

for country in countries:

    values = country.get("values", {})

    has_price = any(
        key.endswith("Price")
        and value is not None
        for key, value in values.items()
    )

    if has_price:
        price_records += 1


quality["priceRecords"] = price_records

atlas["quality"] = quality


# Создаём резервную копию
timestamp = datetime.now().strftime(
    "%Y%m%d_%H%M%S"
)

backup_path = ATLAS_PATH.with_name(
    f"atlas.backup_{timestamp}.json"
)

shutil.copy2(
    ATLAS_PATH,
    backup_path,
)


# Записываем новый atlas.json
with open(
    ATLAS_PATH,
    "w",
    encoding="utf-8"
) as f:

    json.dump(
        atlas,
        f,
        ensure_ascii=False,
        separators=(",", ":"),
    )


print()
print("======================================")
print("Oil Atlas — обновление завершено")
print("======================================")

print(
    f"Стран в Excel:       "
    f"{len(excel_countries)}"
)

print(
    f"Обновлено стран:     "
    f"{updated}"
)

print(
    f"Не сопоставлено:     "
    f"{len(missing)}"
)

print(
    f"НПЗ сохранено:       "
    f"{len(atlas.get('refineries', []))}"
)

print(
    f"Записей с ценами:    "
    f"{price_records}"
)


if missing:

    print()
    print("Не сопоставленные страны:")

    for item in missing:

        print(
            f"  fid={item['fid']} | "
            f"{item['name']}"
        )


print()
print("Резервная копия:")
print(backup_path)

print()
print("Обновлённый файл:")
print(ATLAS_PATH)

print()
print("ГОТОВО")