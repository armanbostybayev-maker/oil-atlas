"""Read-only XLSX adapter. All merging/publication belongs to data/update.mjs."""
import json
import sys
from datetime import date, datetime
import openpyxl

workbook = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
result = {}
try:
    for sheet in sys.argv[2:]:
        if sheet not in workbook.sheetnames:
            raise ValueError(f"Missing worksheet: {sheet}")
        iterator = workbook[sheet].iter_rows(values_only=True)
        headers = next(iterator, None)
        if not headers:
            raise ValueError(f"Empty worksheet: {sheet}")
        required = {"fid"} if sheet == "country_world_globalpetrolprice" else {"Страна", "НПЗ / площадка", "Компания", "Статус источника", "Мощность, барр./сут.", "Мощность, млн т/год", "x", "y"}
        if not required.issubset(headers):
            raise ValueError(f"Missing columns in {sheet}: {sorted(required - set(headers))}")
        rows = []
        for values in iterator:
            if not any(v is not None for v in values):
                continue
            row = {str(k): v.isoformat() if isinstance(v, (date, datetime)) else v for k, v in zip(headers, values) if k is not None}
            if sheet == "country_world_globalpetrolprice" and row.get("fid") is None:
                continue
            rows.append(row)
        result[sheet] = rows
finally:
    workbook.close()
json.dump(result, sys.stdout, ensure_ascii=False, allow_nan=False)
